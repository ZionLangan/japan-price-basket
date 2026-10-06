"""Small public-source probe for the data audit; not the production importer."""

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "audit" / "2026-10-05"
EXTRA_LIBS = ROOT / "tmp" / "audit-python"
if EXTRA_LIBS.exists():
    sys.path.insert(0, str(EXTRA_LIBS))


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.active = None

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.active = {"href": dict(attrs).get("href", ""), "text": ""}

    def handle_data(self, data):
        if self.active is not None:
            self.active["text"] += data

    def handle_endtag(self, tag):
        if tag == "a" and self.active is not None:
            self.active["text"] = re.sub(r"\s+", " ", self.active["text"]).strip()
            self.links.append(self.active)
            self.active = None


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def fetch(name, url, refresh=False):
    OUT.mkdir(parents=True, exist_ok=True)
    info_path = OUT / f"{name}.receipt.json"
    if info_path.exists() and not refresh:
        info = json.loads(info_path.read_text(encoding="utf-8"))
        print(json.dumps({"cached": True, **info}, ensure_ascii=False))
        return
    url = urllib.parse.quote(url, safe=":/?=&%+;,@#")
    req = urllib.request.Request(url, headers={"User-Agent": "JapanInflationDataAudit/0.1", "Accept": "*/*"})
    info = {"name": name, "requested_url": url, "retrieved_at_utc": datetime.now(timezone.utc).isoformat()}
    try:
        with urllib.request.urlopen(req, timeout=40) as response:
            data = response.read(20 * 1024 * 1024 + 1)
            if len(data) > 20 * 1024 * 1024:
                raise ValueError("Audit download exceeds 20 MiB")
            content_type = response.headers.get("Content-Type", "")
            info.update(status=response.status, final_url=response.url, content_type=content_type,
                        content_disposition=response.headers.get("Content-Disposition"),
                        last_modified=response.headers.get("Last-Modified"), etag=response.headers.get("ETag"),
                        bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
            ext = ".bin"
            if data.startswith(b"%PDF"):
                ext = ".pdf"
            elif data.startswith(b"PK\x03\x04"):
                ext = ".xlsx"
            elif data.startswith(b"\xd0\xcf\x11\xe0"):
                ext = ".xls"
            elif "html" in content_type:
                ext = ".html"
            elif "json" in content_type or data.lstrip().startswith((b"{", b"[")):
                ext = ".json"
            elif "csv" in content_type:
                ext = ".csv"
            path = OUT / f"{name}{ext}"
            path.write_bytes(data)
            info["file"] = path.relative_to(ROOT).as_posix()
            if ext == ".html":
                meta_charset = re.search(r"charset\s*=\s*[\"']?\s*([a-zA-Z0-9_-]+)",
                                         data[:10000].decode("ascii", errors="ignore"), re.I)
                charset = response.headers.get_content_charset() or (meta_charset.group(1) if meta_charset else "utf-8")
                text = data.decode(charset, errors="replace")
                parser = Links()
                parser.feed(text)
                links = [{"url": urllib.parse.urljoin(response.url, a["href"]), "text": a["text"]}
                         for a in parser.links if a["href"]]
                write_json(OUT / f"{name}.links.json", links)
                body = re.sub(r"<(script|style)\b[^>]*>.*?</\1>", "", text, flags=re.S | re.I)
                body = unescape(re.sub(r"<[^>]+>", " ", body))
                (OUT / f"{name}.text.txt").write_text(re.sub(r"\s+", " ", body).strip(), encoding="utf-8")
    except (urllib.error.URLError, ValueError) as exc:
        info["error"] = str(exc)
        if isinstance(exc, urllib.error.HTTPError):
            info["status"] = exc.code
    write_json(info_path, info)
    print(json.dumps(info, ensure_ascii=False))


def inspect(name):
    info = json.loads((OUT / f"{name}.receipt.json").read_text(encoding="utf-8"))
    path = ROOT / info["file"]
    result = {"name": name, "file": info["file"]}
    if path.suffix == ".xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        result["sheets"] = []
        for sheet in workbook:
            rows = [{"row": i, "values": list(row)} for i, row in enumerate(sheet.iter_rows(values_only=True), 1)
                    if any(v is not None for v in row)]
            result["sheets"].append({"name": sheet.title, "rows": sheet.max_row,
                                     "columns": sheet.max_column, "nonempty_rows": rows})
        workbook.close()
    elif path.suffix == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(str(path))
        result["sheets"] = [{"name": s.name, "rows": s.nrows, "columns": s.ncols,
                             "nonempty_rows": [{"row": i + 1, "values": s.row_values(i)}
                                               for i in range(s.nrows) if any(v != "" for v in s.row_values(i))]}
                            for s in workbook.sheets()]
    elif path.suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        result["pages"] = [{"page": i + 1, "text": p.extract_text()} for i, p in enumerate(reader.pages)]
    elif path.suffix == ".csv":
        data = path.read_bytes()
        for encoding in ["utf-8-sig", "cp932", "utf-16"]:
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Unknown CSV encoding")
        result.update(encoding=encoding, rows=list(csv.reader(io.StringIO(text))))
    else:
        raise ValueError(f"No tabular inspector for {path.suffix}")
    write_json(OUT / f"{name}.extracted.json", result)
    print(json.dumps({"name": name, "sheets": [{k: v for k, v in s.items() if k != "nonempty_rows"}
                                                for s in result.get("sheets", [])],
                      "pages": len(result.get("pages", [])), "csv_rows": len(result.get("rows", []))},
                     ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["fetch", "inspect"])
    parser.add_argument("name")
    parser.add_argument("url", nargs="?")
    parser.add_argument("--refresh", action="store_true")
    args = parser.parse_args()
    if not re.fullmatch(r"[a-z0-9_-]+", args.name):
        parser.error("Use a simple lowercase audit file name")
    if args.action == "fetch":
        if not args.url:
            parser.error("fetch requires a public URL")
        fetch(args.name, args.url, args.refresh)
    else:
        inspect(args.name)
